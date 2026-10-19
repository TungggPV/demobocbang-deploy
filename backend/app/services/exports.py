"""Asynchronous, ACL-protected transcript exports."""

from __future__ import annotations

import io
import json
import os
import re
import tempfile
from pathlib import Path
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ApiError
from app.models.enums import ExportFormat, JobStatus, MeetingPermission
from app.models.export import Export
from app.models.media import Job
from app.models.meeting import Meeting
from app.models.summary import Summary, SummaryVersion
from app.models.transcript import Speaker, TranscriptVersion
from app.models.translation import Translation, TranslationSegment
from app.models.user import User
from app.providers.storage import MediaStorage, StorageError
from app.schemas.export import ExportCreateRequest, ExportResponse
from app.services.meetings import MeetingService
from app.services.transcript_versions import TranscriptVersionReader

_CONTENT_TYPES = {
    ExportFormat.TXT: "text/plain; charset=utf-8",
    ExportFormat.SRT: "application/x-subrip; charset=utf-8",
    ExportFormat.VTT: "text/vtt; charset=utf-8",
    ExportFormat.JSON: "application/json; charset=utf-8",
    ExportFormat.PDF: "application/pdf",
    ExportFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ExportService:
    def __init__(self, session: AsyncSession, storage: MediaStorage) -> None:
        self._session, self._storage = session, storage
        self._meetings = MeetingService(session)

    async def request(
        self, meeting_id: UUID, user: User, payload: ExportCreateRequest
    ) -> ExportResponse:
        meeting, _ = await self._meetings.require_access(meeting_id, user, MeetingPermission.READ)
        version = await self._current_version(meeting)
        export = Export(
            meeting_id=meeting.id, transcript_version_id=version.id, format=payload.format
        )
        self._session.add(export)
        await self._session.flush()
        job = Job(
            meeting_id=meeting.id,
            job_type="export",
            status=JobStatus.QUEUED,
            progress=0,
            payload={"export_id": str(export.id), "options": payload.options.model_dump()},
        )
        self._session.add(job)
        await self._session.commit()
        await self._session.refresh(export)
        await self._session.refresh(job)
        return self._response(export, job)

    async def get(self, export_id: UUID, user: User) -> ExportResponse:
        export = await self._export_for_read(export_id, user)
        return self._response(export, await self._job(export.id))

    async def download(self, export_id: UUID, user: User) -> tuple[bytes, Export]:
        export = await self._export_for_read(export_id, user)
        job = await self._job(export.id)
        if job is None or job.status != JobStatus.COMPLETED or not export.storage_key:
            raise ApiError("EXPORT_NOT_READY", "The export is not ready to download.", 409)
        try:
            return await self._storage.get_bytes(export.storage_key), export
        except (OSError, StorageError) as exc:
            raise ApiError("EXPORT_STORAGE_ERROR", "The export file is unavailable.", 503) from exc

    async def presigned_download_url(
        self, export_id: UUID, user: User, expires_in: int
    ) -> str | None:
        export = await self._export_for_read(export_id, user)
        job = await self._job(export.id)
        if job is None or job.status != JobStatus.COMPLETED or not export.storage_key:
            raise ApiError("EXPORT_NOT_READY", "The export is not ready to download.", 409)
        disposition = f'attachment; filename="{export.filename or "meeting-export"}"'
        try:
            return await self._storage.presigned_get_url(
                export.storage_key,
                expires_in,
                content_type=export.content_type,
                content_disposition=disposition,
            )
        except StorageError as exc:
            raise ApiError(
                "EXPORT_STORAGE_ERROR", "The export URL could not be created.", 503
            ) from exc

    async def purge_meeting_storage(self, meeting_id: UUID) -> None:
        keys = list(
            await self._session.scalars(
                select(Export.storage_key).where(
                    Export.meeting_id == meeting_id, Export.storage_key.is_not(None)
                )
            )
        )
        try:
            for key in keys:
                await self._storage.delete(key)
        except StorageError as exc:
            raise ApiError(
                "EXPORT_STORAGE_ERROR", "Meeting exports could not be deleted.", 502
            ) from exc

    async def _current_version(self, meeting: Meeting) -> TranscriptVersion:
        if meeting.current_transcript_version_id is None:
            raise ApiError("TRANSCRIPT_NOT_FOUND", "A transcript is required for export.", 422)
        version = await self._session.get(TranscriptVersion, meeting.current_transcript_version_id)
        if version is None:
            raise ApiError("TRANSCRIPT_NOT_FOUND", "A transcript is required for export.", 422)
        return version

    async def _export_for_read(self, export_id: UUID, user: User) -> Export:
        export = await self._session.get(Export, export_id)
        if export is None:
            raise ApiError("EXPORT_NOT_FOUND", "Export was not found.", 404)
        await self._meetings.require_access(export.meeting_id, user, MeetingPermission.READ)
        return export

    async def _job(self, export_id: UUID) -> Job | None:
        jobs = list(await self._session.scalars(select(Job).where(Job.job_type == "export")))
        return next(
            (job for job in jobs if (job.payload or {}).get("export_id") == str(export_id)), None
        )

    @staticmethod
    def _response(export: Export, job: Job | None) -> ExportResponse:
        return ExportResponse(
            id=export.id,
            meeting_id=export.meeting_id,
            transcript_version_id=export.transcript_version_id,
            format=export.format,
            job_id=job.id if job else None,
            status=job.status if job else JobStatus.FAILED,
            filename=export.filename,
            content_type=export.content_type,
            created_at=export.created_at,
        )


class ExportJobRunner:
    """Runs export jobs using the same durable ``jobs`` table as AI jobs."""

    def __init__(self, storage: MediaStorage) -> None:
        self._storage = storage

    async def run(self, session: AsyncSession, job: Job) -> None:
        payload = job.payload or {}
        export = await session.get(Export, UUID(str(payload["export_id"])))
        if export is None:
            raise ValueError("Export request no longer exists.")
        meeting = await session.get(Meeting, export.meeting_id)
        if meeting is None:
            raise ValueError("Meeting no longer exists.")
        version = await session.get(TranscriptVersion, export.transcript_version_id)
        if version is None:
            raise ValueError("Transcript version no longer exists.")
        try:
            segments = await TranscriptVersionReader(session).read(version)
        except ApiError as exc:
            # Missing history is terminal input failure, not a transient provider failure.
            raise ValueError(str(exc)) from exc
        if not segments:
            raise ValueError("Transcript input is empty or invalid.")
        options = {
            "include_timestamps": True,
            "include_speakers": True,
            "include_translation": False,
            "include_summary": False,
            **dict(payload.get("options", {})),
        }
        speakers = {
            item.id: item.name
            for item in await session.scalars(
                select(Speaker).where(Speaker.meeting_id == meeting.id)
            )
        }
        translation = (
            await _translation_by_segment(session, meeting.id, export.transcript_version_id)
            if options["include_translation"]
            else {}
        )
        summary = (
            await _latest_summary(session, meeting.id, export.transcript_version_id)
            if options["include_summary"]
            else None
        )
        document = _document(meeting, export, segments, speakers, translation, summary, options)
        content = _render(export.format, document)
        filename = _filename(meeting.title, export.format)
        key = f"exports/meetings/{meeting.id}/{export.id}.{export.format.value}"
        # Export binaries can be large. Keep the object-storage boundary
        # streaming-friendly and avoid assuming an API/worker shared volume.
        descriptor, temporary_name = tempfile.mkstemp(
            prefix="bocbang-export-", suffix=f".{export.format.value}"
        )
        temporary = Path(temporary_name)
        try:
            with os.fdopen(descriptor, "wb") as output:
                output.write(content)
            await self._storage.put_file(temporary, key)
        finally:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
        export.storage_key, export.filename, export.content_type = (
            key,
            filename,
            _CONTENT_TYPES[export.format],
        )


async def _translation_by_segment(
    session: AsyncSession, meeting_id: UUID, version_id: UUID
) -> dict[UUID, str]:
    translation = await session.scalar(
        select(Translation)
        .where(
            Translation.meeting_id == meeting_id, Translation.transcript_version_id == version_id
        )
        .order_by(Translation.created_at.desc())
        .limit(1)
    )
    if translation is None:
        return {}
    rows = await session.scalars(
        select(TranslationSegment).where(TranslationSegment.translation_id == translation.id)
    )
    return {row.logical_segment_id: row.text for row in rows}


async def _latest_summary(
    session: AsyncSession, meeting_id: UUID, version_id: UUID
) -> dict[str, object] | None:
    summary = await session.scalar(
        select(Summary)
        .where(Summary.meeting_id == meeting_id, Summary.transcript_version_id == version_id)
        .order_by(Summary.created_at.desc())
        .limit(1)
    )
    if summary is None:
        return None
    row = await session.scalar(
        select(SummaryVersion)
        .where(SummaryVersion.summary_id == summary.id)
        .order_by(SummaryVersion.version_no.desc())
        .limit(1)
    )
    return row.content if row else None


def _document(
    meeting: Meeting,
    export: Export,
    segments: list[dict[str, object]],
    speakers: dict[UUID, str],
    translation: dict[UUID, str],
    summary: dict[str, object] | None,
    options: dict[str, object],
) -> dict[str, object]:
    data = {
        "meeting": {
            "title": meeting.title,
            "description": meeting.description,
            "language_mode": meeting.language_mode,
            "created_at": meeting.created_at.isoformat(),
            "transcript_version_id": str(export.transcript_version_id),
        },
        "options": options,
        "segments": [],
    }
    for item in segments:
        row: dict[str, object] = {
            "id": str(item["logical_segment_id"]),
            "start_ms": item["start_ms"],
            "end_ms": item["end_ms"],
            "text": item["text"],
        }
        if options["include_speakers"]:
            row["speaker"] = speakers.get(item["speaker_id"]) if item["speaker_id"] else None
        if options["include_translation"] and item["logical_segment_id"] in translation:
            row["translation"] = translation[item["logical_segment_id"]]
        data["segments"].append(row)
    if options["include_summary"] and summary is not None:
        data["summary"] = summary
    return data


def _render(format: ExportFormat, data: dict[str, object]) -> bytes:
    if format == ExportFormat.JSON:
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
    if format in (ExportFormat.SRT, ExportFormat.VTT):
        return _subtitles(data, format).encode("utf-8")
    if format == ExportFormat.TXT:
        return _plain_text(data).encode("utf-8")
    if format == ExportFormat.DOCX:
        return _render_docx(data)
    if format == ExportFormat.PDF:
        return _render_pdf(data)
    raise ValueError(f"Unsupported format: {format}")


def _render_pdf(data: dict[str, object]) -> bytes:
    import html
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from reportlab.platypus import (
        HRFlowable,
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    reg_font = next(
        (
            p
            for p in [
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                "C:/Windows/Fonts/arial.ttf",
                "C:/Windows/Fonts/segoeui.ttf",
            ]
            if os.path.exists(p)
        ),
        None,
    )
    bold_font = next(
        (
            p
            for p in [
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
                "C:/Windows/Fonts/arialbd.ttf",
                "C:/Windows/Fonts/segoeuib.ttf",
            ]
            if os.path.exists(p)
        ),
        None,
    )
    italic_font = next(
        (
            p
            for p in [
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
                "C:/Windows/Fonts/ariali.ttf",
                "C:/Windows/Fonts/segoeuii.ttf",
            ]
            if os.path.exists(p)
        ),
        None,
    )

    if reg_font is None:
        raise ValueError("A Unicode PDF font is not configured.")

    font_name = "ExportDocFont"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, reg_font))
        pdfmetrics.registerFont(TTFont(f"{font_name}-Bold", bold_font or reg_font))
        pdfmetrics.registerFont(TTFont(f"{font_name}-Italic", italic_font or reg_font))
        registerFontFamily(
            font_name,
            normal=font_name,
            bold=f"{font_name}-Bold",
            italic=f"{font_name}-Italic",
            boldItalic=f"{font_name}-Bold",
        )

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_page_decorations(num_pages)
                super().showPage()
            super().save()

        def draw_page_decorations(self, page_count):
            self.saveState()
            self.setFont(font_name, 8)
            self.setFillColor(colors.HexColor("#64748b"))
            self.setStrokeColor(colors.HexColor("#e2e8f0"))
            self.setLineWidth(0.5)
            self.line(40, 36, 555, 36)
            self.drawString(40, 24, "Bóc Băng AI — Hệ thống bóc băng & ghi âm thông minh")
            self.drawRightString(555, 24, f"Trang {self._pageNumber} / {page_count}")
            self.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=36, bottomMargin=50
    )

    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle(
        "DocBrand",
        parent=styles["Normal"],
        fontName=f"{font_name}-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#4f46e5"),
    )
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName=f"{font_name}-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
    )
    meta_lbl = ParagraphStyle(
        "MetaLbl",
        parent=styles["Normal"],
        fontName=f"{font_name}-Bold",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#64748b"),
    )
    meta_val = ParagraphStyle(
        "MetaVal",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#0f172a"),
    )
    sec_style = ParagraphStyle(
        "DocSec",
        parent=styles["Normal"],
        fontName=f"{font_name}-Bold",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#1e293b"),
    )
    speaker_style = ParagraphStyle(
        "DocSpk",
        parent=styles["Normal"],
        fontName=f"{font_name}-Bold",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#2563eb"),
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9.5,
        leading=14.5,
        textColor=colors.HexColor("#334155"),
    )
    trans_style = ParagraphStyle(
        "DocTrans",
        parent=styles["Normal"],
        fontName=f"{font_name}-Italic",
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor("#475569"),
    )
    summary_body = ParagraphStyle(
        "SumBody",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=13.5,
        textColor=colors.HexColor("#1e3a5f"),
    )

    story = []
    meeting = data["meeting"]
    title = str(meeting.get("title") or "Bản ghi cuộc họp")
    lang_mode = str(meeting.get("language_mode") or "auto")
    lang_label = {"vi": "Tiếng Việt", "en": "Tiếng Anh", "auto": "Tự động nhận diện"}.get(
        lang_mode, lang_mode
    )

    story.append(Paragraph("BÓC BĂNG AI &nbsp;·&nbsp; TÀI LIỆU XUẤT BẢN GHI CUỘC HỌP", brand_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"<b>{html.escape(title)}</b>", title_style))
    story.append(Spacer(1, 10))

    unique_speakers = []
    for item in data.get("segments", []):
        spk = item.get("speaker")
        if spk and spk not in unique_speakers:
            unique_speakers.append(spk)
    speakers_str = ", ".join(unique_speakers) if unique_speakers else "Tự động phân tách"
    created_at = (
        meeting.get("created_at", "")[:16].replace("T", " ") if meeting.get("created_at") else ""
    )

    meta_data = [
        [
            Paragraph("<b>Thời gian:</b>", meta_lbl),
            Paragraph(html.escape(created_at) if created_at else "—", meta_val),
            Paragraph("<b>Ngôn ngữ:</b>", meta_lbl),
            Paragraph(html.escape(lang_label), meta_val),
        ],
        [
            Paragraph("<b>Số đoạn:</b>", meta_lbl),
            Paragraph(f"{len(data.get('segments', []))} đoạn", meta_val),
            Paragraph("<b>Người nói:</b>", meta_lbl),
            Paragraph(html.escape(speakers_str), meta_val),
        ],
    ]
    meta_table = Table(meta_data, colWidths=[70, 185, 75, 185])
    meta_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
                ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#e2e8f0")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#f1f5f9")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(meta_table)
    story.append(Spacer(1, 14))

    if data.get("summary"):
        summary_content = data["summary"]
        story.append(Paragraph("<b>TÓM TẮT CUỘC HỌP</b>", sec_style))
        story.append(Spacer(1, 4))
        story.append(
            HRFlowable(
                width="100%",
                thickness=0.5,
                color=colors.HexColor("#e2e8f0"),
                spaceBefore=2,
                spaceAfter=8,
            )
        )

        summary_paragraphs = []
        if isinstance(summary_content, dict):
            for k, v in summary_content.items():
                k_label = {
                    "summary": "Tóm tắt",
                    "overview": "Tổng quan",
                    "key_points": "Các ý chính",
                    "action_items": "Hành động tiếp theo",
                    "decisions": "Quyết định",
                }.get(k, k.capitalize())
                if isinstance(v, list):
                    summary_paragraphs.append(Paragraph(f"<b>{html.escape(k_label)}:</b>", meta_lbl))
                    for sub in v:
                        summary_paragraphs.append(
                            Paragraph(f"• {html.escape(str(sub))}", summary_body)
                        )
                elif isinstance(v, str):
                    summary_paragraphs.append(
                        Paragraph(f"<b>{html.escape(k_label)}:</b> {html.escape(v)}", summary_body)
                    )
        elif isinstance(summary_content, str):
            summary_paragraphs.append(Paragraph(html.escape(summary_content), summary_body))

        if summary_paragraphs:
            sum_table = Table([[p] for p in summary_paragraphs], colWidths=[515])
            sum_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f0fdf4")),
                        ("BOX", (0, 0), (-1, -1), 0.75, colors.HexColor("#bbf7d0")),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ("LEFTPADDING", (0, 0), (-1, -1), 10),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                    ]
                )
            )
            story.append(sum_table)
            story.append(Spacer(1, 14))

    story.append(Paragraph("<b>NỘI DUNG CHI TIẾT BẢN GHI</b>", sec_style))
    story.append(Spacer(1, 4))
    story.append(
        HRFlowable(
            width="100%",
            thickness=0.75,
            color=colors.HexColor("#e2e8f0"),
            spaceBefore=2,
            spaceAfter=10,
        )
    )

    include_timestamps = data.get("options", {}).get("include_timestamps", True)
    include_speakers = data.get("options", {}).get("include_speakers", True)

    for item in data.get("segments", []):
        time_part = ""
        if include_timestamps:
            t_start = _format_human_time(item["start_ms"])
            t_end = _format_human_time(item["end_ms"])
            time_part = f'<font color="#64748b" size="8">[{t_start} - {t_end}]</font>'

        spk_part = ""
        if include_speakers and item.get("speaker"):
            spk_part = f'<b>{html.escape(str(item["speaker"]))}</b>'
        elif include_speakers:
            spk_part = '<font color="#94a3b8">Chưa xác định</font>'

        header_html = f"{spk_part} &nbsp; {time_part}".strip()
        text_html = html.escape(str(item.get("text", ""))).replace("\n", "<br/>")

        block = []
        if header_html:
            block.append(Paragraph(header_html, speaker_style))
            block.append(Spacer(1, 3))
        block.append(Paragraph(text_html, body_style))

        if item.get("translation"):
            trans_html = html.escape(str(item["translation"])).replace("\n", "<br/>")
            block.append(Spacer(1, 3))
            block.append(Paragraph(f"<b>Bản dịch:</b> {trans_html}", trans_style))

        block.append(Spacer(1, 8))
        story.append(KeepTogether(block))

    doc.build(story, canvasmaker=NumberedCanvas)
    return buf.getvalue()


def _render_docx(data: dict[str, object]) -> bytes:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_brand = doc.add_paragraph()
    r_brand = p_brand.add_run("BÓC BĂNG AI · TÀI LIỆU XUẤT BẢN GHI CUỘC HỌP")
    r_brand.font.size = Pt(8.5)
    r_brand.font.bold = True
    r_brand.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)

    meeting = data["meeting"]
    title = str(meeting.get("title") or "Bản ghi cuộc họp")
    h = doc.add_heading(level=0)
    run_h = h.add_run(title)
    run_h.font.size = Pt(18)
    run_h.font.bold = True
    run_h.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    lang_mode = str(meeting.get("language_mode") or "auto")
    lang_label = {"vi": "Tiếng Việt", "en": "Tiếng Anh", "auto": "Tự động nhận diện"}.get(
        lang_mode, lang_mode
    )
    created_at = (
        meeting.get("created_at", "")[:16].replace("T", " ") if meeting.get("created_at") else ""
    )

    table = doc.add_table(rows=2, cols=4)
    table.autofit = False
    metadata = [
        [("Thời gian:", True), (created_at or "—", False), ("Ngôn ngữ:", True), (lang_label, False)],
        [
            ("Số đoạn:", True),
            (f"{len(data.get('segments', []))} đoạn", False),
            ("Người nói:", True),
            ("Tự động phân tách", False),
        ],
    ]
    for row_idx, row_data in enumerate(metadata):
        row = table.rows[row_idx]
        for col_idx, (text, is_bold) in enumerate(row_data):
            cell = row.cells[col_idx]
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
            cell._tc.get_or_add_tcPr().append(shd)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9)
            run.font.bold = is_bold
            if is_bold:
                run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            else:
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph()

    if data.get("summary"):
        h_sum = doc.add_heading(level=1)
        r_sum = h_sum.add_run("TÓM TẮT CUỘC HỌP")
        r_sum.font.size = Pt(13)
        r_sum.font.bold = True
        r_sum.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        summary_content = data["summary"]
        if isinstance(summary_content, dict):
            for k, v in summary_content.items():
                k_label = {
                    "summary": "Tóm tắt",
                    "overview": "Tổng quan",
                    "key_points": "Các ý chính",
                    "action_items": "Hành động tiếp theo",
                    "decisions": "Quyết định",
                }.get(k, k.capitalize())
                p_k = doc.add_paragraph()
                r_k = p_k.add_run(f"{k_label}:")
                r_k.font.bold = True
                if isinstance(v, list):
                    for sub in v:
                        p_sub = doc.add_paragraph(style="List Bullet")
                        p_sub.add_run(str(sub))
                elif isinstance(v, str):
                    p_k.add_run(f" {v}")
        elif isinstance(summary_content, str):
            doc.add_paragraph(summary_content)

        doc.add_paragraph()

    h_sec = doc.add_heading(level=1)
    run_sec = h_sec.add_run("NỘI DUNG CHI TIẾT BẢN GHI")
    run_sec.font.size = Pt(13)
    run_sec.font.bold = True
    run_sec.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    include_timestamps = data.get("options", {}).get("include_timestamps", True)
    include_speakers = data.get("options", {}).get("include_speakers", True)

    for item in data.get("segments", []):
        time_part = ""
        if include_timestamps:
            t_start = _format_human_time(item["start_ms"])
            t_end = _format_human_time(item["end_ms"])
            time_part = f"[{t_start} - {t_end}]"

        spk_part = ""
        if include_speakers and item.get("speaker"):
            spk_part = str(item["speaker"])
        elif include_speakers:
            spk_part = "Chưa xác định"

        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_before = Pt(6)
        p_meta.paragraph_format.space_after = Pt(2)
        if spk_part:
            r_spk = p_meta.add_run(spk_part)
            r_spk.font.bold = True
            r_spk.font.size = Pt(10)
            r_spk.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        if time_part:
            r_time = p_meta.add_run(f"  {time_part}")
            r_time.font.size = Pt(9)
            r_time.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        p_text = doc.add_paragraph()
        p_text.paragraph_format.space_before = Pt(0)
        p_text.paragraph_format.space_after = Pt(4)
        r_text = p_text.add_run(str(item.get("text", "")))
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        if item.get("translation"):
            p_trans = doc.add_paragraph()
            p_trans.paragraph_format.space_before = Pt(0)
            p_trans.paragraph_format.space_after = Pt(4)
            r_tl = p_trans.add_run("Bản dịch: ")
            r_tl.font.bold = True
            r_tl.font.size = Pt(9)
            r_tl.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
            r_tv = p_trans.add_run(str(item["translation"]))
            r_tv.font.italic = True
            r_tv.font.size = Pt(9)
            r_tv.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _plain_text(data: dict[str, object]) -> str:
    meeting = data["meeting"]
    title = str(meeting.get("title") or "Bản ghi cuộc họp")
    lang_mode = str(meeting.get("language_mode") or "auto")
    created_at = (
        meeting.get("created_at", "")[:16].replace("T", " ") if meeting.get("created_at") else ""
    )

    lines = [
        "=" * 50,
        f"BẢN GHI CUỘC HỌP: {title}",
        f"Thời gian: {created_at} | Ngôn ngữ: {lang_mode}",
        "=" * 50,
        "",
    ]
    if meeting.get("description"):
        lines.extend([str(meeting["description"]), ""])

    if data.get("summary"):
        lines.extend(["--- TÓM TẮT CUỘC HỌP ---"])
        summary_content = data["summary"]
        if isinstance(summary_content, dict):
            for k, v in summary_content.items():
                if isinstance(v, list):
                    lines.append(f"{k}:")
                    for sub in v:
                        lines.append(f"  • {sub}")
                else:
                    lines.append(f"{k}: {v}")
        else:
            lines.append(str(summary_content))
        lines.append("")

    lines.append("--- NỘI DUNG CHI TIẾT ---")
    include_timestamps = data.get("options", {}).get("include_timestamps", True)
    include_speakers = data.get("options", {}).get("include_speakers", True)

    for item in data.get("segments", []):
        prefix = ""
        if include_timestamps:
            prefix += f"[{_format_human_time(item['start_ms'])} - {_format_human_time(item['end_ms'])}] "
        if include_speakers and item.get("speaker"):
            prefix += f"{item['speaker']}: "
        lines.append(prefix + str(item.get("text", "")))
        if item.get("translation"):
            lines.append(f"  Bản dịch: {item['translation']}")
    return "\n".join(lines) + "\n"


def _subtitles(data: dict[str, object], format: ExportFormat) -> str:
    lines = ["WEBVTT", ""] if format == ExportFormat.VTT else []
    for index, item in enumerate(data.get("segments", []), start=1):
        if format == ExportFormat.SRT:
            lines.append(str(index))
        separator = "." if format == ExportFormat.VTT else ","
        lines.append(
            f"{_timestamp(item['start_ms'], separator)} --> {_timestamp(item['end_ms'], separator)}"
        )
        label = f"{item['speaker']}: " if item.get("speaker") else ""
        lines.append(label + str(item.get("text", "")))
        lines.append("")
    return "\n".join(lines)


def _timestamp(value: object, separator: str = ".") -> str:
    ms = int(value)
    hours, ms = divmod(ms, 3_600_000)
    minutes, ms = divmod(ms, 60_000)
    seconds, ms = divmod(ms, 1000)
    return f"{hours:02}:{minutes:02}:{seconds:02}{separator}{ms:03}"


def _format_human_time(value: object) -> str:
    ms = int(value)
    hours, ms = divmod(ms, 3_600_000)
    minutes, ms = divmod(ms, 60_000)
    seconds, _ = divmod(ms, 1000)
    if hours > 0:
        return f"{hours:02}:{minutes:02}:{seconds:02}"
    return f"{minutes:02}:{seconds:02}"


def _filename(title: str, format: ExportFormat) -> str:
    stem = re.sub(r"[^\w.-]+", "-", title, flags=re.UNICODE).strip(".-") or "meeting"
    return f"{stem}.{format.value}"
